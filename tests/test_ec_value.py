"""Value 3BR EC mode: helpers, plus a full run against a fake PropertyGuru on localhost."""
import json
import threading
from datetime import date
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

from pgscraper import config, ec_value
from pgscraper.parse import extract_bedrooms, extract_photos

JPEG = b"\xff\xd8\xff\xe0" + b"0" * 64


def test_extract_bedrooms():
    assert extract_bedrooms("3 Beds | 2 Baths | 1,023 sqft | 3 Bedrooms") == 3
    assert extract_bedrooms("Spacious 3+1 bedrooms") == 3
    assert extract_bedrooms("4BR corner") == 4
    assert extract_bedrooms("no rooms mentioned") is None


def test_extract_photos_keeps_biggest_and_skips_agents_and_other_listings():
    html = ('"https://sg1-p.pgimgs.com/listing/500111/UPHO.1.V550/a.jpg" '
            '"https://sg1-p.pgimgs.com/listing/500111/UPHO.1.V800/a.jpg" '
            '"https://sg1-p.pgimgs.com/listing/500111/UPHO.2.V800/b.jpg" '
            '"https://sg1-p.pgimgs.com/agent/9/AGPHO.9.V120/me.jpg" '
            '"https://sg1-p.pgimgs.com/listing/500999/UPHO.7.V800/other.jpg"')
    assert extract_photos(html, "500111") == [
        "https://sg1-p.pgimgs.com/listing/500111/UPHO.1.V800/a.jpg",
        "https://sg1-p.pgimgs.com/listing/500111/UPHO.2.V800/b.jpg",
    ]


def test_eligible_projects_need_mop_and_top_2016_or_later():
    today = date(2026, 9, 30)
    ok = ec_value.eligible_projects(today)
    assert "Sea Horizon" in ok and "Rivercove Residences" in ok
    assert "Piermont Grand" not in ok  # still within MOP
    assert "The Topiary" in ok and "Waterbay" in ok  # TOP'd early 2016
    assert "Waterwoods" not in ok  # TOP'd Dec 2015, not in the list


# ---------------------------------------------------------------- end to end

def page(listings):
    data = {"props": {"pageProps": {"listings": listings}}}
    return (f'<html><head><title>PG</title></head><body>'
            f'<script id="__NEXT_DATA__" type="application/json">{json.dumps(data)}</script></body></html>')


def item(i, price, sqft, title="The Criterion"):
    return {"id": i, "url": f"/listing/for-sale-the-criterion-{i}", "title": title,
            "address": "7 Yishun Street 51", "price": {"value": price}, "floorArea": {"value": sqft}}


def detail(i, price, sqft, beds):
    photos = " ".join(f'<img src="SERVER/pgimgs.com/listing/{i}/UPHO.{k}.V800/p{k}.jpg">' for k in (1, 2))
    return (f"<html><body>{page([item(i, price, sqft)])} Listing {i}. {beds} Beds, 2 Baths. "
            f"High Floor. {photos}</body></html>")


FIRST = [item(500001, 1_450_000, 1023), item(500002, 1_900_000, 1098), item(500003, 1_200_000, 893)]
PAGES = {
    # the project's own page, then a keyword search whose page 1 repeats it and page 2 has a new one
    "/property-for-sale/at-the-criterion-22281": page(FIRST),
    "/property-for-sale/at-the-criterion-22281/2": page([]),
    "/property-for-sale": page(FIRST),
    "/property-for-sale/2": page([item(500007, 1_550_000, 1033)]),
    "/property-for-sale/3": page([]),
    "/listing/for-sale-the-criterion-500001": detail(500001, 1_450_000, 1023, 3),
    "/listing/for-sale-the-criterion-500007": detail(500007, 1_550_000, 1033, 3),
    "/listing/for-sale-the-criterion-500004": detail(500004, 1_500_000, 1098, 4),
    "/listing/for-sale-the-criterion-500005": "<html><body>This listing is no longer available</body></html>",
}


class H(BaseHTTPRequestHandler):
    def do_GET(self):
        path = self.path.split("?")[0]
        if path.endswith(".jpg"):
            self.send_response(200)
            self.send_header("Content-Type", "image/jpeg")
            self.end_headers()
            self.wfile.write(JPEG)
            return
        body = PAGES.get(path)
        self.send_response(200 if body else 404)
        self.send_header("Content-Type", "text/html")
        self.end_headers()
        host = f"http://{self.headers['Host']}"
        self.wfile.write((body or "not found").replace("SERVER", host).encode())

    def log_message(self, *a):
        pass


@pytest.fixture
def server():
    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def test_ec_value_end_to_end(server, tmp_path, monkeypatch):
    shortlist = tmp_path / "shortlist.csv"
    shortlist.write_text(
        "project,listing_id,price,size_sqft,block_address,check,url\n"
        f"The Criterion,500004,1500000,1098,,,{server}/listing/for-sale-the-criterion-500004\n"
        f"The Criterion,500005,1400000,1033,,,{server}/listing/for-sale-the-criterion-500005\n")
    monkeypatch.setattr(config, "EC_PROJECTS",
                        {"The Criterion": ("2018-02-26", "Yishun Street 51", "Yishun", "the-criterion-22281")})
    monkeypatch.setattr(config, "EC_PROJECT_URL", server + "/property-for-sale/at-{slug}")
    monkeypatch.setattr(config, "EC_SHORTLIST_CSV", str(shortlist))
    monkeypatch.setattr(config, "EC_SEARCH_URL", server + "/property-for-sale?freetext={q}")
    monkeypatch.setattr(config, "OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setattr(config, "PAGE_DELAY_SECONDS", (0, 0))
    monkeypatch.setattr("pgscraper.parse.BASE_URL", server)

    assert ec_value.main(["--contact", "June · 9123 4567"]) == 0
    [out] = list((tmp_path / "out").iterdir())
    deck = (out / "deck.html").read_text()
    # 500001 and 500007 fit; 500002 over budget, 500003 too small, 500004 is a 4-bed, 500005 is gone
    assert "500001" in deck and "500007" in deck
    assert not any(x in deck for x in ("500002", "500003", "500004", "500005"))
    assert deck.count("data:image/jpeg;base64,") == 4
    assert "June · 9123 4567" in deck
    assert sorted(p.name for p in (out / "photos").iterdir()) == [
        "500001_1.jpg", "500001_2.jpg", "500004_1.jpg", "500004_2.jpg", "500007_1.jpg", "500007_2.jpg"]
    from openpyxl import load_workbook
    cov = list(load_workbook(out / "shortlist.xlsx")["Coverage"].iter_rows(min_row=2, values_only=True))
    assert [r[:5] for r in cov] == [("The Criterion", "2018-02-26", 4, 2, 2)]


def test_project_of_ignores_punctuation():
    from pgscraper.parse import Listing
    l = Listing("1", url="https://x/listing/for-sale-citylife-tampines-1", title="CityLife @ Tampines")
    assert ec_value.project_of(l, ["CityLife @ Tampines", "The Vales"]) == "CityLife @ Tampines"
    l2 = Listing("2", title="Northwave EC Executive Condominium")
    assert ec_value.project_of(l2, ["Northwave"]) == "Northwave"
