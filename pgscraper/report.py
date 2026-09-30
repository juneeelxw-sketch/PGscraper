"""Write the Excel workbook and a flat CSV of the current listings."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .history import Change
from .parse import Listing
from .rental import IDEAL, grade

MONEY = '"S$"#,##0'
MONEY_SIGNED = '"+S$"#,##0;"-S$"#,##0;0'
SQFT = "#,##0"
PSF = '"S$"#,##0.00'

STATUS_FILL = {
    "NEW": "C6EFCE",
    "RELISTED": "C6EFCE",
    "PRICE DOWN": "FFEB9C",
    "PRICE UP": "FFEB9C",
    "REMOVED": "FFC7CE",
}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], formats: dict[str, str],
           link_col: str | None = "URL"):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        if h in formats:
            for cell in ws[letter][1:]:
                cell.number_format = formats[h]
        if h == link_col:
            for cell in ws[letter][1:]:
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows]) + 2
        ws.column_dimensions[letter].width = min(width, 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _color_status(ws, col: int = 1) -> None:
    for row in ws.iter_rows(min_row=2):
        fill = STATUS_FILL.get(row[col - 1].value)
        if fill:
            for c in row:
                c.fill = PatternFill("solid", fgColor=fill)


def _rental_sheet(wb: Workbook, listings: list[Listing], by_id: dict[str, Change]) -> None:
    rows = []
    for l in listings:
        fit, note = grade(l)
        ch = by_id.get(l.listing_id)
        rows.append([
            fit, ch.status if ch else "", l.title, l.price, l.size_sqft, l.psf, l.bedrooms, l.bathrooms,
            l.property_type, l.mrt, note, ch.prev_price if ch and ch.status.startswith("PRICE") else None,
            ch.first_seen if ch else "", l.address, l.agent, l.listed_date, l.listing_id, l.url,
        ])
    rows.sort(key=lambda r: (r[0] != IDEAL, r[3] or 0))
    ws = _sheet(wb, "Listings",
                ["Fit", "Status", "Development", "Rent", "Size (sqft)", "PSF", "Beds", "Baths", "Type",
                 "Nearest MRT", "Why not ideal", "Prev Rent", "First Seen", "Address", "Agent", "Listed",
                 "Listing ID", "URL"],
                rows, {"Rent": MONEY, "Prev Rent": MONEY, "Size (sqft)": SQFT, "PSF": PSF})
    _color_status(ws, col=2)
    for row in ws.iter_rows(min_row=2):
        if row[0].value == IDEAL:
            row[0].fill = PatternFill("solid", fgColor=STATUS_FILL["NEW"])


def write_xlsx(path: Path, listings: list[Listing], changes: list[Change], history_rows: list[dict],
               run_at: str, rental: bool = False) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    by_id = {c.listing_id: c for c in changes}

    if rental:
        _rental_sheet(wb, listings, by_id)
    else:
        _sale_sheets(wb, listings, by_id)

    # What changed since last run
    order = ["NEW", "RELISTED", "PRICE DOWN", "PRICE UP", "REMOVED"]
    rows = [[c.status, c.unit_no or "(not stated)", c.prev_price, c.price, c.price_change, c.size_sqft,
             c.first_seen, c.title, c.agent, c.url]
            for c in sorted(changes, key=lambda c: order.index(c.status) if c.status in order else 99)
            if c.status != "UNCHANGED"]
    ws = _sheet(wb, "Changes",
                ["Status", "Unit No.", "Prev Price", "Price", "Change", "Size (sqft)", "First Seen",
                 "Title", "Agent", "URL"],
                rows, {"Prev Price": MONEY, "Price": MONEY, "Change": MONEY_SIGNED, "Size (sqft)": SQFT})
    _color_status(ws)

    # Full log, including this run
    all_rows = history_rows + [{"run_at": run_at, **l.to_row()} for l in listings]
    rows = [[r["run_at"], r["unit_no"], _f(r["price"]), _f(r["size_sqft"]), _f(r["psf"]), r["agent"],
             r["listing_id"], r["url"]] for r in all_rows]
    _sheet(wb, "Price History", ["Run At", "Unit No.", "Price", "Size (sqft)", "PSF", "Agent", "Listing ID",
                                 "URL"],
           rows, {"Price": MONEY, "Size (sqft)": SQFT, "PSF": PSF})

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _sale_sheets(wb: Workbook, listings: list[Listing], by_id: dict[str, Change]) -> None:
    # Current listings, cheapest psf first
    rows = []
    for l in sorted(listings, key=lambda l: (l.psf is None, l.psf or 0)):
        ch = by_id.get(l.listing_id)
        rows.append([
            ch.status if ch else "", l.unit_no or "(not stated)", l.floor_hint, l.price, l.size_sqft, l.psf,
            ch.prev_price if ch and ch.status.startswith("PRICE") else None,
            ch.first_seen if ch else "", l.title, l.agent, l.agency, l.listed_date, l.listing_id, l.url,
        ])
    ws = _sheet(wb, "Listings",
                ["Status", "Unit No.", "Floor Hint", "Price", "Size (sqft)", "PSF", "Prev Price",
                 "First Seen", "Title", "Agent", "Agency", "Listed", "Listing ID", "URL"],
                rows, {"Price": MONEY, "Prev Price": MONEY, "Size (sqft)": SQFT, "PSF": PSF})
    _color_status(ws)

    # One line per unit (the same unit is often listed by several agents)
    groups: dict[str, list[Listing]] = defaultdict(list)
    for l in listings:
        for u in (l.unit_no.split(", ") if l.unit_no else ["(not stated)"]):
            groups[u].append(l)
    rows = []
    for unit, ls in sorted(groups.items()):
        prices = [l.price for l in ls if l.price]
        sizes = sorted({l.size_sqft for l in ls if l.size_sqft})
        low = min(prices) if prices else None
        rows.append([unit, len(ls), low, max(prices) if prices else None,
                     ", ".join(f"{s:,.0f}" for s in sizes),
                     round(low / sizes[0], 2) if low and sizes else None,
                     ", ".join(sorted({l.agent for l in ls if l.agent}))])
    _sheet(wb, "By Unit", ["Unit No.", "# Listings", "Lowest Price", "Highest Price", "Size(s) sqft",
                           "Lowest PSF", "Agents"],
           rows, {"Lowest Price": MONEY, "Highest Price": MONEY, "Lowest PSF": PSF}, link_col=None)


def _f(v):
    try:
        return float(v) if v not in ("", None) else None
    except (TypeError, ValueError):
        return None


def write_csv(path: Path, listings: list[Listing], changes: list[Change]) -> None:
    status = {c.listing_id: c.status for c in changes}
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["status", "unit_no", "floor_hint", "price", "size_sqft", "psf", "title", "agent",
                    "listing_id", "url", "bedrooms", "bathrooms", "property_type", "mrt"])
        for l in listings:
            w.writerow([status.get(l.listing_id, ""), l.unit_no, l.floor_hint, l.price, l.size_sqft, l.psf,
                        l.title, l.agent, l.listing_id, l.url, l.bedrooms, l.bathrooms, l.property_type,
                        l.mrt])
