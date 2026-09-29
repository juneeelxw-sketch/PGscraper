"""Runs the real CLI against a fake PropertyGuru served from localhost."""
import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

from pgscraper import __main__ as cli, config


def page(listings):
    data = {"props": {"pageProps": {"listings": listings}}}
    return (f'<html><head><title>PG</title></head><body><h1>Results</h1>'
            f'<script id="__NEXT_DATA__" type="application/json">{json.dumps(data)}</script></body></html>')


def item(i, price, title="Ubi Techpark", address="10 Ubi Crescent"):
    return {"id": i, "url": f"/listing/for-sale-{i}", "title": title, "address": address,
            "price": {"value": price}, "floorArea": {"value": 1200}}


PAGES = {
    "/property-for-sale": page([item(111111, 1_200_000), item(222222, 1_300_000, "Elsewhere", "5 Kaki Bukit Rd")]),
    "/property-for-sale/2": page([item(333333, 900_000)]),
    "/property-for-sale/3": page([]),
    "/listing/for-sale-111111": "<html><body>Unit #05-21 for sale, high floor</body></html>",
    "/listing/for-sale-333333": "<html><body>Level 2 corner unit, ramp-up</body></html>",
}


class H(BaseHTTPRequestHandler):
    def do_GET(self):
        body = PAGES.get(self.path.split("?")[0])
        self.send_response(200 if body else 404)
        self.send_header("Content-Type", "text/html")
        self.end_headers()
        self.wfile.write((body or "not found").encode())

    def log_message(self, *a):
        pass


@pytest.fixture
def server():
    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def test_cli_end_to_end(server, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "HISTORY_CSV", str(tmp_path / "obs.csv"))
    monkeypatch.setattr(config, "OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setattr(config, "PAGE_DELAY_SECONDS", (0, 0))
    monkeypatch.setattr("pgscraper.parse.BASE_URL", server)

    assert cli.main(["--url", f"{server}/property-for-sale?freetext=Ubi"]) == 0
    csv_text = (tmp_path / "out" / "latest.csv").read_text()
    assert "#05-21" in csv_text and "Level 2" in csv_text and "Elsewhere" not in csv_text
    assert len(list((tmp_path / "out").glob("*.xlsx"))) == 1

    # second run: listing 333333 gone, 111111 cheaper
    PAGES["/property-for-sale"] = page([item(111111, 1_150_000)])
    PAGES["/property-for-sale/2"] = page([])
    assert cli.main(["--url", f"{server}/property-for-sale?freetext=Ubi"]) == 0
    from openpyxl import load_workbook
    wb = load_workbook(sorted((tmp_path / "out").glob("*.xlsx"))[-1])
    statuses = [r[0] for r in wb["Changes"].iter_rows(min_row=2, values_only=True)]
    assert statuses == ["PRICE DOWN", "REMOVED"]
